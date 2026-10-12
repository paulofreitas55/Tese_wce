#!/usr/bin/env python3
"""
Passo 0.3 — Mapa de rótulos clínicos (não toca no dataset).

Cada seed tem uma descrição clínica em texto livre, tirada do nome do ficheiro imagemN_<descrição>
(ex.: "Angiectasias_punctiformes_do_jejuno"). Este passo põe cada descrição distinta numa de 14
categorias e diz se é lesão (1) ou não (0). Faz-se em duas vezes:

  1) PROPOSTA — regras de palavras propõem a categoria de cada descrição e geram a folha de revisão:
         python scripts/05_rotulos.py
     Depois revês à mão docs/label_map_revisao.xlsx (coluna "categoria"; as linhas amarelas primeiro).

  2) APLICAR — depois da revisão: valida a folha, grava o mapa final e acrescenta as colunas
     categoria e lesao às tabelas das janelas e dos frames:
         python scripts/05_rotulos.py --aplicar

Seeds com mais de uma descrição: se só uma for patologia, fica essa; duas patologias diferentes ficam
por resolver à mão. Para impor a categoria de uma janela concreta: docs/rotulos_janelas_manual.csv
(colunas split,caso,seed,categoria,motivo) — sobrepõe-se a tudo.

Lê:    data/derived/manifest.csv · data/derived/janelas_v1.csv · (2) data/derived/annotations_v1.csv
Gera:  (1) docs/label_map_revisao.xlsx · docs/criterios_categorias.md (só se ainda não existir)
           outputs/05_rotulos_proposta.txt   ← é este que se lê
       (2) data/derived/label_map.csv · janelas_v1.csv e annotations_v1.csv (+ categoria, lesao)
           outputs/05_rotulos_resumo.txt     ← é este que se lê
"""
import argparse
import datetime
import os
import re
import shutil
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import pandas as pd

K3 = ["split", "caso", "seed"]
SO_EXTENSAO = {"", "png", "jpg", "jpeg"}
COLS_NOVAS_JAN = ["categoria", "lesao", "categoria_origem"]
COLS_NOVAS_FR = ["categoria", "lesao"]

# ------------------------------------------------------------------------------------------------
# As 14 categorias. "termos" são procurados no texto já normalizado (minúsculas, sem acentos,
# "_" e pontuação → espaço). A ordem da lista é a ordem das linhas nas tabelas.
# ------------------------------------------------------------------------------------------------
CATEGORIAS = [
    dict(cod="angiectasia", lesao=1, nome="Angiectasia",
         entra="Lesão vascular: angiectasia/angioectasia, angiodisplasia, telangiectasia, lesão vascular, "
               "GAVE (ectasia vascular antral), incluindo gralhas (angeictasia, anfiectasia). Não inclui "
               "linfangiectasia nem flebectasia.",
         termos=[r"angi\w{0,3}ect", r"angeict", r"anfiect", r"angiodisplas", r"angiodysplas", r"vascular(es)?\b",
                 r"\bgave\b"],
         legivel="angiect…, angioect…, angeict…, anfiect…, angiodisplas…, vascular(es), GAVE"),
    dict(cod="erosao", lesao=1, nome="Erosão",
         entra="Erosão, erosões, microerosões, mucosa erosiva. Aftas/lesões aftoides entram aqui, "
               "marcadas para confirmar.",
         termos=[r"eros", r"\baft[ao]", r"aphth"],
         legivel="eros…, afta/aftoide/aftosa"),
    dict(cod="ulcera", lesao=1, nome="Úlcera",
         entra="Úlcera, ulceração, lesão ulcerada.",
         termos=[r"ulcer"],
         legivel="ulcer…"),
    dict(cod="polipo", lesao=1, nome="Pólipo",
         entra="Pólipo, lesão polipoide, polipose, pseudopólipo.",
         termos=[r"polip", r"polyp"],
         legivel="polip…"),
    dict(cod="eritema", lesao=1, nome="Eritema",
         entra="Eritema, mucosa eritematosa, hiperemia, enantema.",
         termos=[r"eritem", r"erythem", r"hiperem", r"hyperem", r"enantem"],
         legivel="eritem…, hiperem…, enantem…"),
    dict(cod="xantelasma", lesao=1, nome="Xantelasma",
         entra="Xantelasma, xantoma.",
         termos=[r"xantel", r"xantom", r"xanth"],
         legivel="xantel…, xantom…"),
    dict(cod="linfangiectasia", lesao=1, nome="Linfangiectasia",
         entra="Linfangiectasia (dilatação linfática: pontos ou vilosidades esbranquiçadas). Categoria "
               "própria: o texto contém 'angiect' e, no trabalho anterior, era contada como angiectasia.",
         termos=[r"linf\w*ect", r"lymphangiect"],
         legivel="linf…ect… (linfangiectasia e gralhas)"),
    dict(cod="sangue", lesao=1, nome="Sangue",
         entra="Sangue, hemorragia, sangramento, coágulo, hematina, vestígios hemáticos, melena.",
         termos=[r"sangu", r"sangr", r"hemorr", r"haemorr", r"hemat", r"coagul", r"melen", r"bleed", r"blood"],
         legivel="sangu…, sangr…, hemorr…, hemat…, coagul…, melen…"),
    dict(cod="hiperplasia_nodular", lesao=1, nome="Hiperplasia nodular linfoide",
         entra="Hiperplasia nodular linfoide (HNL), nódulos ou folículos linfoides, mucosa nodular.",
         termos=[r"hiperplas", r"hyperplas", r"n(od|do)ul", r"\bhnl\b", r"linfoid", r"lymphoid", r"folicul"],
         legivel="hiperplas…, nodul…, HNL, linfoid…, folicul…"),
    dict(cod="estenose", lesao=1, nome="Estenose",
         entra="Estenose, estreitamento ou redução do lúmen, anel, diafragma, segmento pouco distensível.",
         termos=[r"estenos", r"stenos", r"estreit", r"strictur", r"\banel\b", r"diafragma", r"distensiv",
                 r"reducao (do |de )?(lumen|luminal|calibre)"],
         legivel="estenos…, estreit…, anel, diafragma, distensív…, redução do lúmen/luminal/de calibre"),
    dict(cod="subepitelial", lesao=1, nome="Lesão subepitelial / massa",
         entra="Abaulamento, protrusão, lesão subepitelial ou submucosa, massa, tumor, neoplasia, "
               "lesão vegetante.",
         termos=[r"subepitel", r"subepithel", r"submucos", r"abaul", r"protus", r"protrus", r"\bmassa\b",
                 r"tumor", r"tumour", r"tumefac", r"neopla", r"vegetant", r"lipoma", r"\bgist\b"],
         legivel="subepitel…, submucos…, abaul…, protus…/protrus…, massa, tumor…, neopla…, vegetant…"),
    dict(cod="outra_patologia", lesao=1, nome="Outra patologia",
         entra="Outros achados patológicos: edema, congestão, fissura, friabilidade, divertículo, atrofia, "
               "cicatriz, exsudado/fibrina, flebectasia/varizes, tom azulado/violáceo, inflamação ('-ite': "
               "bulbite, duodenite, ileíte…), enteropatia/gastropatia, scalloping, apagamento ou convergência de "
               "pregas, deformação, área deprimida, aspeto invaginado, escara de laqueação, parasitas. Também "
               "'lesão'/'achado' sem tipo (ex.: Lesao_no_ileon).",
         termos=[r"edem", r"congest", r"fissur", r"friab", r"friav", r"divertic", r"atrofi", r"atroph",
                 r"cicatri", r"exsud", r"fibrin", r"flebect", r"phlebect", r"variz", r"varic", r"inflam",
                 r"espessa", r"(gastr|duoden|bulb|jejun|ile|col|enter|esofag|proct)ite\b", r"invagin",
                 r"scallop", r"azulad", r"violac", r"apagament", r"convergencia", r"deforma", r"deprimid",
                 r"escara", r"laqueac", r"(entero|gastro|colo|bulbo)pat", r"duplo lumen", r"ascari", r"oxiur",
                 r"parasit", r"helmint", r"\bvermes?\b"],
         legivel="edem…, congest…, fissur…, friab…, divertic…, atrofi…, cicatri…, exsud…, fibrin…, "
                 "flebect…, variz…, inflam…, espessa…, …ite (gastrite, bulbite…), …patia, scalloping, azulad…, "
                 "violác…, apagamento, convergência, deform…, deprimid…, invagin…, escara, laqueação, ascaris, "
                 "oxiúro, parasit…; sem outro termo: "
                 "lesão, achado, anomalia, alteração, suspeita"),
    dict(cod="nao_lesao", lesao=0, nome="Não-lesão (landmark)",
         entra="Sem patologia no texto: marcos anatómicos e localização (primeira imagem gástrica, "
               "duodenal ou cecal, piloro, papila, válvula ileocecal, antro, incisura, vértice bulbar, D2…), "
               "pregas, mucosa normal, conteúdo ou bílis, tatuagem, clip, cicatriz cirúrgica, ostomia, "
               "corpo estranho.",
         termos=[r"primeir", r"\bfirst\b", r"\bultim", r"\blast\b", r"pilor", r"pylor", r"valvul", r"\bvalv",
                 r"ileocec", r"ileo cec", r"\bvic\b", r"papil", r"ampol", r"vater", r"linha z", r"\bz line",
                 r"juncao", r"\bcardia\b", r"esofag", r"estomag", r"gastric", r"\bgastr", r"duoden", r"bulbo",
                 r"jejun", r"\bile[oa]?\b", r"\bileal\b", r"\bileon", r"\bileum", r"\bcego\b", r"\bceco\b",
                 r"cecal", r"\bca?ecum\b", r"colon", r"\bcol(o|ic[ao]s?)\b", r"apendic", r"treitz", r"\bret(o|al)\b",
                 r"\brect(o|al)\b", r"sigmoid", r"tatua", r"\bclip", r"agraf", r"sutur", r"anastomos", r"\bbili",
                 r"\bmuco\b", r"conteud", r"\brestos\b", r"aliment", r"\bnormal", r"\bantr(o|al)\b", r"incisura",
                 r"recesso", r"vertice", r"\bbulbar\b", r"\bbolbo\b", r"\bd[1-4]\b", r"ostomi", r"\bestoma\b",
                 r"\bbolsa\b", r"corpo estranho", r"cirurg", r"angulac", r"\bpregas?\b"],
         legivel="primeira/first, última/last, piloro, válvula, ileocecal, papila, linha Z, cárdia, esófago, "
                 "estômago, duodeno, bulbo, jejuno, íleo/íleon, cego/cecal, cólon, apêndice, reto, tatuagem, "
                 "clip, agrafe, sutura, anastomose, bílis, muco, conteúdo, restos alimentares, normal, antro, "
                 "incisura, recesso, vértice, bulbar, D1–D4, ostomia, bolsa, corpo estranho, cirúrgic…, "
                 "angulação, prega(s)"),
    dict(cod="qualidade", lesao=0, nome="Qualidade",
         entra="O texto diz que não se consegue avaliar: má preparação ou limpeza, resíduos, bolhas ou "
               "espuma, observação prejudicada ou inconclusiva.",
         termos=[r"prepara", r"prejudic", r"inconclusiv", r"impureza", r"espuma", r"bolha", r"residu", r"\bturv",
                 r"fecal", r"\bfezes", r"limpeza", r"debris", r"bubbl", r"\bfoam", r"nao avaliav"],
         legivel="prepara…, prejudic…, inconclusiv…, impureza…, espuma, bolha…, resíduo…, turv…, "
                 "fezes/fecal, limpeza, não avaliável"),
]
CODIGOS = [c["cod"] for c in CATEGORIAS]
LESAO = {c["cod"]: c["lesao"] for c in CATEGORIAS}
TERMOS = {c["cod"]: [re.compile(p) for p in c["termos"]] for c in CATEGORIAS}
TIPADAS = CODIGOS[:11]                       # patologias com tipo definido (todas menos outra_patologia)
PRINCIPAIS = ["angiectasia", "erosao", "ulcera", "polipo", "eritema", "xantelasma"]
ORDEM = {c: i for i, c in enumerate([""] + CODIGOS)}

GENERICOS = [re.compile(p) for p in [r"\bles(ao|oes)\b", r"\blesion", r"achado", r"anomal", r"anormal", r"abnormal",
                                     r"alterac", r"\bfinding", r"suspeit", r"patolog"]]
NEGACAO = re.compile(r"\b(sem|nao|ausencia de|without) [a-z0-9]+")
NAO = re.compile(r"\bnao\b")
INCERTEZA = re.compile(r"esboc|possiv|possib|provav|suspeit|duvid|eventual|sugestiv|incert|esclarec|nao se exclu")
AFTA = re.compile(r"\baft[ao]|aphth")
CONFIRMAR = [(re.compile(p), m) for p, m in [
    (r"corpo estranho", "corpo estranho: ficou não-lesão (não é lesão da mucosa) — confirmar"),
    (r"recesso", "recesso: anatomia ou divertículo? confirmar"),
    (r"ascari|oxiur|parasit|helmint|\bvermes?\b", "parasita: ficou outra_patologia — confirmar"),
    (r"\bgave\b", "GAVE (ectasia vascular antral): ficou angiectasia — confirmar"),
    (r"\banel\b|diafragma", "anel/diafragma: estenose? confirmar"),
    (r"distensiv", "pouco distensível: estenose? confirmar"),
    (r"angulac", "angulação: anatomia? confirmar"),
    (r"proeminent", "proeminente: normal ou espessado? confirmar"),
    (r"azulad|violac", "tom azulado/violáceo: lesão vascular? confirmar"),
    (r"invagin|duplo lumen|deprimid", "achado estrutural pouco claro: confirmar"),
    (r"escara|laqueac", "escara de laqueação: confirmar a categoria"),
    (r"cicatri", "cicatriz: conta como lesão? confirmar"),
    (r"ostomi|\bestoma\b|\bbolsa\b", "ostomia/bolsa (cirurgia): ficou não-lesão — confirmar"),
]]
VAZIAS = set("""de do da dos das no na nos nas em e a o os as com ao aos um uma of the and in at on image imagem
 imagens mucosa area zona nivel porcao segmento regiao parte inicio fim terco proximal distal medio media inicial
 final terminal superior inferior angulo flexura hepatica esplenica fundo corpo lumen luz parede aspeto aspecto
 visao panoramica tempo transito passagem entrada saida chegada transicao ver observacao visualizacao exame
 avaliacao ma deficiente insuficiente fraca boa razoavel limitada dificil dificultada pobre poor bad""".split())
PALAVRA_LINF = re.compile(r"[a-z0-9]*(linf[a-z0-9]*ect|lymphangiect)[a-z0-9]*")
SINONIMOS = {"landmark": "nao_lesao", "nao_lesao_landmark": "nao_lesao", "hnl": "hiperplasia_nodular",
             "hiperplasia_nodular_linfoide": "hiperplasia_nodular", "outras_patologias": "outra_patologia",
             "outra": "outra_patologia", "lesao_subepitelial": "subepitelial", "massa": "subepitelial",
             "hemorragia": "sangue", "erosoes": "erosao", "ulceras": "ulcera", "polipos": "polipo",
             "angiectasias": "angiectasia", "linfangiectasias": "linfangiectasia"}
SIM = {"sim", "s", "x", "1", "yes", "y", "ok", "true", "verdadeiro"}

# Regra do script antigo (build_dataset.py), copiada tal e qual, só para comparar contagens.
PAT_KW_ANTIGO = ["angiect", "angioect", "angeict", "anfiect", "angiodisplas", "flebect",
                 "erosa", "eroso", "microeros", "ulcer", "polip", "xantel", "eritema",
                 "eritemat", "linfangiect", "hiperplasia", "nodular", "hnl", "abaulament",
                 "protus", "hemorrag", "sangue", "coagulo", "hematic", "hematin",
                 "subepitelial", "estenos", "congesta", "congestiv", "fissura", "edema",
                 "friabil", "friavel", "divertic", "tumor", "neoplas", "vegetant",
                 "lesao", "lesoes", "achado", "anomal", "suspeit"]
QUAL_KW_ANTIGO = ["ma preparac", "deficiente preparac", "observacao prejud",
                  "observacao inconclusiva", "impurezas", "espuma"]


def rotulo_antigo(desc):
    d = re.sub(r"\s+", " ", str(desc).lower().strip().replace("_", " "))
    if any(k in d for k in PAT_KW_ANTIGO):
        return "lesao"
    if any(k in d for k in QUAL_KW_ANTIGO):
        return "qualidade"
    return "outro"


# ------------------------------------------------------------------------------------------------
def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def limpar_desc(d):                          # igual ao 02_tabela.py
    d = str(d).strip(" .")
    return "" if d.lower() in SO_EXTENSAO else d


def em_branco(regex, t):                     # apaga sem mudar as posições do resto do texto
    return regex.sub(lambda m: " " * len(m.group()), t)


def palavra(t, s, e):
    while s > 0 and t[s - 1].isalnum():
        s -= 1
    while e < len(t) and t[e].isalnum():
        e += 1
    return t[s:e]


def procurar(padroes, t):
    return sorted({(m.start(), palavra(t, m.start(), m.end())) for p in padroes for m in p.finditer(t)})


def classificar(desc):
    """Proposta automática para uma descrição: (categoria, termos encontrados, motivos para rever)."""
    t = norm(desc)
    base = em_branco(NEGACAO, t)                       # "sem sangue", "sem alterações" não contam
    achados = {}
    linf = procurar(TERMOS["linfangiectasia"], base)
    if linf:
        achados["linfangiectasia"] = linf
        base = em_branco(PALAVRA_LINF, base)           # "linfangiectasia" contém "angiect"
    for c in TIPADAS + ["outra_patologia"]:
        if c != "linfangiectasia":
            r = procurar(TERMOS[c], base)
            if r:
                achados[c] = r
    gen = procurar(GENERICOS, base)
    qual = procurar(TERMOS["qualidade"], base)
    nl = procurar(TERMOS["nao_lesao"], base)
    negacao = bool(NEGACAO.search(t) or NAO.search(t))

    motivos = []
    tip = {c: v for c, v in achados.items() if c in TIPADAS}
    if tip:
        por_ordem = sorted(tip, key=lambda c: tip[c][0][0])
        cat = por_ordem[0]                             # o nome da lesão costuma vir primeiro
        if len(tip) > 1:
            motivos.append("várias categorias: " + ", ".join(por_ordem) + f" — ficou {cat} (aparece primeiro)")
        if "outra_patologia" in achados:
            motivos.append("também tem termo de outra_patologia (" +
                           ", ".join(w for _, w in achados["outra_patologia"]) + ")")
    elif "outra_patologia" in achados:
        cat = "outra_patologia"
    elif gen:
        cat = "outra_patologia"
        motivos.append("só termo genérico (" + ", ".join(w for _, w in gen) + ")")
    elif qual:
        cat = "qualidade"
    elif nl or negacao:
        cat = "nao_lesao"
    else:
        cat = ""
        motivos.append("nenhuma regra: escolher à mão")

    if LESAO.get(cat) == 1 and qual:
        motivos.append("patologia e qualidade no mesmo texto")
    if INCERTEZA.search(t) or "?" in str(desc):
        motivos.append("texto com dúvida")
    if negacao:
        motivos.append("texto com negação")
    if cat == "erosao" and AFTA.search(t) and not re.search(r"eros", base):
        motivos.append("afta: ficou em erosão — confirmar")
    if cat == "angiectasia" and all(w.startswith("vascul") for _, w in achados["angiectasia"]):
        motivos.append("'vascular' sem 'angiectasia': confirmar")
    if cat == "hiperplasia_nodular" and not re.search(r"hiperplas|hyperplas|\bhnl\b|linfoid|lymphoid|folicul", t):
        motivos.append("nódulo sem 'hiperplasia/linfoide': confirmar (pode ser pólipo ou subepitelial)")
    if cat == "subepitelial" and re.search(r"massa|tumor|tumour|neopla|vegetant", t):
        motivos.append("massa/tumor: confirmar")
    if cat == "outra_patologia" and re.search(r"flebect|phlebect|variz|varic", t):
        motivos.append("vascular venosa (flebectasia/variz): confirmar")
    if cat == "outra_patologia" and re.search(r"cirurg", t) and \
            all(w.startswith("cicatri") for _, w in achados.get("outra_patologia", [])):
        cat = "nao_lesao"                              # cicatriz cirúrgica: não é lesão
    for rx, msg in CONFIRMAR:
        if rx.search(t):
            motivos.append(msg)
    if cat in ("nao_lesao", "qualidade"):              # rede de segurança: palavras que nenhuma regra explica
        conhecidas = {x for _, w in qual + nl + [y for v in achados.values() for y in v] for x in w.split()}
        sobra = [w for w in base.split() if w not in VAZIAS and w not in conhecidas and not w.isdigit()]
        if sobra:
            motivos.append("palavras sem regra: " + ", ".join(dict.fromkeys(sobra)) + " — confirmar que não é lesão")
    antigo = rotulo_antigo(desc)
    if cat and LESAO[cat] != (antigo == "lesao"):
        motivos.append(f"a regra antiga dizia '{antigo}'")

    partes = [f"{c}: " + ", ".join(dict.fromkeys(w for _, w in v)) for c, v in achados.items()]
    for nome, r in [("genérico", gen), ("qualidade", qual), ("localização/landmark", nl)]:
        if r:
            partes.append(f"{nome}: " + ", ".join(dict.fromkeys(w for _, w in r)))
    return cat, " | ".join(partes), "; ".join(motivos), antigo


# ------------------------------------------------------------------------------------------------
def ler_janelas(caminho):
    if not Path(caminho).is_file():
        sys.exit(f"Não encontrei {caminho}: corre primeiro o 02_tabela.py e o 04_grupos.py.")
    jan = pd.read_csv(caminho, dtype=str, keep_default_na=False)
    for c in ["n_box", "seed_tem_box", "excluida", "n_descricoes"]:
        jan[c] = pd.to_numeric(jan[c], errors="coerce").fillna(0).astype(int)
    jan["descs"] = [[d for d in s.split(" | ") if d] if s else [] for s in jan["descricao"]]
    return jan


def descricoes_manifest(caminho):
    """Descrições limpas (como no 02_tabela.py) e quantas seriam à moda do script antigo (sem limpar o fim)."""
    if not Path(caminho).is_file():
        sys.exit(f"Não encontrei {caminho}: corre primeiro o 01_inventario.py.")
    m = pd.read_csv(caminho, dtype=str, keep_default_na=False, usecols=["caminho", "tipo", "descricao"])
    m = m[m["tipo"] == "relatorio_descricao"]
    limpas = sorted({x for x in m["descricao"].map(limpar_desc) if x})
    antigas = set()
    for c in m["caminho"]:
        mm = re.match(r"imagem\d+_(.+)", Path(c).name, re.I)
        if mm and "videoframe" not in mm.group(1).lower():
            d = mm.group(1).rsplit(".", 1)[0].strip()
            if d:
                antigas.add(d)
    return limpas, len(antigas)


def contagens_por_descricao(jan):
    """Para cada descrição: nº de janelas (não excluídas) que a usam, por split e com caixa."""
    linhas = {}
    us = jan[jan["excluida"] == 0]
    for r in us.itertuples(index=False):
        for d in r.descs:
            c = linhas.setdefault(d, dict(n_janelas=0, n_treino=0, n_teste=0, n_com_caixa=0, n_caixa_seed=0,
                                          exemplo="", _ex_box=False))
            c["n_janelas"] += 1
            c["n_treino"] += int(r.split.endswith("Treino"))
            c["n_teste"] += int(r.split.endswith("Teste"))
            c["n_com_caixa"] += int(r.n_box > 0)
            c["n_caixa_seed"] += int(r.seed_tem_box == 1)
            if not c["exemplo"] or (r.seed_tem_box == 1 and not c["_ex_box"]):
                c["exemplo"] = f"{r.split}/{r.caso}/{r.seed}"
                c["_ex_box"] = r.seed_tem_box == 1
    return linhas


def categoria_janela(descs, mapa):
    """(categoria, lesao, origem) de uma janela a partir das categorias das suas descrições."""
    if not descs:
        return "", "", "sem_descricao"
    cats = [mapa.get(d, "") for d in descs]
    if any(c not in LESAO for c in cats):
        return "", "", "descricao_sem_categoria"
    if len(set(cats)) == 1:
        return cats[0], str(LESAO[cats[0]]), "descricao"
    pato = sorted({c for c in cats if LESAO[c] == 1})
    if len(pato) == 1:
        return pato[0], "1", "varias_descricoes"
    if len(pato) > 1:
        return "", "", "varias_patologias"
    c = "qualidade" if "qualidade" in cats else "nao_lesao"
    return c, "0", "varias_descricoes"


def ler_manual(caminho):
    if not Path(caminho).is_file():
        return {}
    df = pd.read_csv(caminho, dtype=str, keep_default_na=False)
    erros, man = [], {}
    for i, r in df.iterrows():
        c = cat_norm(r.get("categoria", ""))
        if c not in LESAO:
            erros.append(f"  {caminho}, linha {i + 2}: categoria '{r.get('categoria', '')}' desconhecida")
            continue
        man[(r["split"], str(r["caso"]).zfill(4), str(r["seed"]))] = (c, r.get("motivo", ""))
    if erros:
        sys.exit("Erros no ficheiro de categorias manuais por janela:\n" + "\n".join(erros))
    return man


def aplicar_janelas(jan, mapa, manual):
    res = [categoria_janela(d, mapa) for d in jan["descs"]]
    jan["categoria"] = [r[0] for r in res]
    jan["lesao"] = [r[1] for r in res]
    jan["categoria_origem"] = [r[2] for r in res]
    usados = 0
    for i, k in enumerate(zip(jan["split"], jan["caso"], jan["seed"])):
        if k in manual:
            c = manual[k][0]
            jan.iat[i, jan.columns.get_loc("categoria")] = c
            jan.iat[i, jan.columns.get_loc("lesao")] = str(LESAO[c])
            jan.iat[i, jan.columns.get_loc("categoria_origem")] = "manual"
            usados += 1
    return usados


def cat_norm(v):
    c = norm(v).replace(" ", "_")
    return SINONIMOS.get(c, c)


# ------------------------------------------------------------------------------------------------
def bloco_tabela(L, jan, mapa_desc):
    """Tabela por categoria (a 'Tabela 1'): descrições e janelas por split, com caixa e com caixa no seed."""
    us = jan[jan["excluida"] == 0]
    n_desc = Counter(mapa_desc.values())
    L.append(f"  {'categoria':<20}{'lesão':>6}{'descr.':>8}{'janelas':>9}{'treino':>8}{'teste':>7}"
             f"{'c/caixa':>9}{'c/caixa seed':>14}")
    for c in CODIGOS + [""]:
        g = us[us["categoria"] == c]
        if c == "" and len(g) == 0 and n_desc.get("", 0) == 0:
            continue
        L.append(f"  {(c or '(sem categoria)'):<20}{(str(LESAO[c]) if c else '?'):>6}{n_desc.get(c, 0):>8}"
                 f"{len(g):>9}{int(g['split'].str.endswith('Treino').sum()):>8}"
                 f"{int(g['split'].str.endswith('Teste').sum()):>7}{int((g['n_box'] > 0).sum()):>9}"
                 f"{int(g['seed_tem_box'].sum()):>14}")
    for nome, sel in [("lesão (1)", us["lesao"] == "1"), ("não-lesão (0)", us["lesao"] == "0"),
                      ("por resolver", us["lesao"] == "")]:
        g = us[sel]
        if nome == "por resolver" and len(g) == 0:
            continue
        L.append(f"  {nome:<20}{'':>6}{'':>8}{len(g):>9}{int(g['split'].str.endswith('Treino').sum()):>8}"
                 f"{int(g['split'].str.endswith('Teste').sum()):>7}{int((g['n_box'] > 0).sum()):>9}"
                 f"{int(g['seed_tem_box'].sum()):>14}")
    L.append(f"  (janelas não excluídas: {len(us)}; uma janela com duas descrições conta uma vez, na categoria da janela)")


def bloco_conjuntos(L, jan, lesao_col="lesao", titulo=""):
    us = jan[jan["excluida"] == 0]
    les = us[lesao_col] == "1"
    nles = us[lesao_col] == "0"
    sb = us["seed_tem_box"] == 1
    cx = us["n_box"] > 0
    te = us["split"].str.endswith("Teste")
    tr = us["split"].str.endswith("Treino")
    sem = us[les & ~sb]
    if titulo:
        L.append(f"  {titulo}")
    L.append(f"  janelas com caixa no seed: {int(sb.sum())} = lesão {int((les & sb).sum())} + não-lesão "
             f"{int((nles & sb).sum())}" + (f" + por resolver {int((~les & ~nles & sb).sum())}"
                                            if (~les & ~nles & sb).any() else ""))
    L.append(f"  ALL-SEEDBOX (lesão com caixa no seed): {int((les & sb).sum())}   (diagnóstico antigo: 1 346)")
    L.append(f"  patologia sem caixa no seed: {len(sem)} = sem caixa nenhuma {int((sem['n_box'] == 0).sum())}"
             f" + caixa só nos vizinhos {int((sem['n_box'] > 0).sum())}   (diagnóstico antigo: 605 = 560 + 45)")
    L.append(f"  não-lesão com caixa: {int((nles & cx).sum())} (com caixa no seed: {int((nles & sb).sum())})"
             f"   (diagnóstico antigo: 35)")
    L.append(f"  TEST-BOX de lesão: {int((les & cx & te).sum())}   (todas as janelas de teste com caixa: "
             f"{int((cx & te).sum())}; diagnóstico antigo: 459)")
    L.append(f"  TRAIN-BOX de lesão: {int((les & cx & tr).sum())}   (todas as janelas de treino com caixa: "
             f"{int((cx & tr).sum())}; diagnóstico antigo: 986)")


def bloco_classes(L, jan):
    us = jan[(jan["excluida"] == 0) & (jan["lesao"] == "1")]
    classe = us["categoria"].where(us["categoria"].isin(PRINCIPAIS), "outras patologias")
    L.append(f"  {'classe':<20}{'janelas tr/te':>15}{'c/ caixa no seed tr/te':>25}")
    for c in PRINCIPAIS + ["outras patologias"]:
        g = us[classe == c]
        t, s = g["split"].str.endswith("Treino"), g["split"].str.endswith("Teste")
        L.append(f"  {c:<20}{int(t.sum()):>8} / {int(s.sum()):<5}{int((t & (g['seed_tem_box'] == 1)).sum()):>15}"
                 f" / {int((s & (g['seed_tem_box'] == 1)).sum())}")


def bloco_multiplas(L, jan, mapa):
    multi = jan[jan["descs"].map(len) >= 2]
    L.append(f"  {len(multi)} janelas")
    for r in multi.itertuples(index=False):
        cats = " | ".join(f"{d} → {mapa.get(d, '') or '?'}" for d in r.descs)
        L.append(f"    {r.split}/{r.caso}/{r.seed} (caixa no seed: {'sim' if r.seed_tem_box else 'não'}; caixas na "
                 f"janela: {r.n_box}): {cats}")
        L.append(f"      → janela fica: {r.categoria or 'POR RESOLVER'} ({r.categoria_origem})"
                 + ("; sem caixa no seed, não entra em ALL-SEEDBOX, CAL nem VAL-EX" if not r.seed_tem_box else ""))


def escrever(L, caminho):
    texto = "\n".join(L)
    print(texto)
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    Path(caminho).write_text(texto + "\n", encoding="utf-8")
    print(f"\nResumo gravado em {caminho}")


def gravar_csv(df, caminho):
    tmp = f"{caminho}.tmp"
    df.to_csv(tmp, index=False)
    os.replace(tmp, caminho)                  # só substitui quando a escrita terminou


# ------------------------------------------------------------------------------------------------
INSTRUCOES = [
    "1. Para cada linha, lê a descrição (coluna B) e confirma a categoria proposta (coluna C).",
    "   Se estiver errada, escolhe outra na lista que aparece ao clicar na célula.",
    "2. Começa pelas linhas amarelas (rever = sim): o motivo está na coluna H. Para as ver só a elas,",
    "   usa o filtro da coluna G.",
    "3. Depois de ver uma linha, escreve sim na coluna D (revisto). Podes parar e continuar noutro dia.",
    "4. Coluna E (notas): só quando mudas a categoria ou ficas em dúvida.",
    "5. Não alteres as colunas A e B: são a chave que liga a linha ao dataset.",
    "6. Duas lesões no mesmo texto: fica a do nome da lesão (o substantivo), não a do adjetivo",
    "   (\"Polipo_ulcerado\" → polipo). Em dúvida se é lesão: outra_patologia + nota.",
    "7. A folha \"categorias\" diz o que entra em cada categoria.",
    "8. Quando acabares, grava como .xlsx com o mesmo nome e volta a pô-la em docs/.",
]


def gravar_folha(df, caminho):
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        sys.exit("Falta o openpyxl neste ambiente: pip install openpyxl  (e volta a correr o script)")
    cats = pd.DataFrame([{"categoria": c["cod"], "lesao": c["lesao"], "nome": c["nome"],
                          "o que entra": c["entra"], "palavras procuradas": c["legivel"]} for c in CATEGORIAS])
    with pd.ExcelWriter(caminho, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="label_map", index=False)
        cats.to_excel(xw, sheet_name="categorias", index=False)
        pd.DataFrame({"como rever": INSTRUCOES}).to_excel(xw, sheet_name="como_rever", index=False)
        ws = xw.sheets["label_map"]
        n, ncol = len(df), len(df.columns)
        amarelo = PatternFill("solid", fgColor="FFF2CC")
        verde = PatternFill("solid", fgColor="D9EAD3")
        col = {c: i + 1 for i, c in enumerate(df.columns)}
        for j in range(1, ncol + 1):
            ws.cell(1, j).font = Font(bold=True)
        for c in ["categoria", "revisto", "notas"]:
            ws.cell(1, col[c]).fill = verde
        for i, rv in enumerate(df["rever"], start=2):
            if rv == "sim":
                for j in range(1, ncol + 1):
                    ws.cell(i, j).fill = amarelo
        larguras = {"id": 8, "descricao": 46, "categoria": 21, "revisto": 9, "notas": 30, "categoria_auto": 20,
                    "rever": 7, "motivo": 55, "termos": 50, "regra_antiga": 12, "exemplo": 28}
        for c, j in col.items():
            ws.column_dimensions[get_column_letter(j)].width = larguras.get(c, 10)
        for i in range(2, n + 2):
            ws.cell(i, col["motivo"]).alignment = Alignment(wrap_text=False)
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = f"A1:{get_column_letter(ncol)}{n + 1}"
        lc = get_column_letter(col["categoria"])
        dv = DataValidation(type="list", formula1='"' + ",".join(CODIGOS) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Categoria inválida",
                            error="Escolhe uma das 14 categorias da lista.")
        ws.add_data_validation(dv)
        dv.add(f"{lc}2:{lc}{n + 1}")
        lr = get_column_letter(col["revisto"])
        dv2 = DataValidation(type="list", formula1='"sim"', allow_blank=True)
        ws.add_data_validation(dv2)
        dv2.add(f"{lr}2:{lr}{n + 1}")
        wc = xw.sheets["categorias"]
        for j, w in enumerate([22, 7, 30, 90, 70], start=1):
            wc.column_dimensions[get_column_letter(j)].width = w
            wc.cell(1, j).font = Font(bold=True)
        xw.sheets["como_rever"].column_dimensions["A"].width = 110


def ler_revisao(caminho):
    p = Path(caminho)
    if not p.is_file():
        sys.exit(f"Não encontrei {caminho}. Corre primeiro sem --aplicar (gera a folha) e revê-a.")
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            df = pd.read_excel(p, sheet_name="label_map", dtype=str)
        except ValueError:
            df = pd.read_excel(p, sheet_name=0, dtype=str)
    else:
        df = None
        for enc in ("utf-8-sig", "cp1252"):
            try:
                df = pd.read_csv(p, dtype=str, keep_default_na=False, sep=None, engine="python", encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        if df is None:
            sys.exit(f"Não consegui ler {caminho}.")
    df = df.fillna("")
    df.columns = [norm(c).replace(" ", "_") for c in df.columns]
    for c in df.columns:
        df[c] = df[c].astype(str).str.strip()
    return df


def criterios_md(caminho):
    if Path(caminho).exists():
        return False
    L = ["# Critérios das categorias clínicas (passo 0.3)", "",
         f"Rascunho gerado por scripts/05_rotulos.py em {datetime.date.today():%d/%m/%Y}. Podes editar à "
         "vontade: o script não volta a escrever este ficheiro.", "",
         "Cada seed tem uma descrição clínica em texto livre (nome do ficheiro imagemN_<descrição>). Cada "
         "descrição distinta foi posta numa de 14 categorias: primeiro por regras de palavras (texto em "
         "minúsculas, sem acentos), depois revista à mão, linha a linha (docs/label_map_revisao.xlsx). "
         "O mapa final está em data/derived/label_map.csv.", "",
         "## Categorias", "",
         "| # | Categoria | Lesão | O que entra | Palavras procuradas |",
         "|---|---|---|---|---|"]
    for i, c in enumerate(CATEGORIAS, 1):
        L.append(f"| {i} | {c['nome']} (`{c['cod']}`) | {c['lesao']} | {c['entra']} | {c['legivel']} |")
    L += ["", "## Regras de decisão", "",
          "1. Termos de mais de um tipo de patologia: fica o tipo cujo termo aparece primeiro no texto (o nome "
          "da lesão costuma vir antes do adjetivo: \"Polipo_ulcerado\" → pólipo); a linha é revista à mão.",
          "2. `outra_patologia` só quando não há termo de um tipo definido. \"Lesão\"/\"achado\" sem tipo também "
          "vão para aqui (ex.: Lesao_no_ileon).",
          "3. Termos de localização (\"jejuno\", \"íleon\"…) só decidem quando não há patologia no texto: "
          "nesse caso a descrição é não-lesão.",
          "4. Sem patologia, qualidade prevalece sobre não-lesão.",
          "5. \"sem X\" não conta (ex.: \"Erosao_sem_sangue\" é erosão).",
          "6. Texto com dúvida (\"possível\", \"suspeita\") fica na categoria da lesão e é revisto à mão.",
          "7. Seeds com mais de uma descrição: se só uma é patologia, fica essa; duas patologias diferentes são "
          "decididas à mão e registadas em docs/rotulos_janelas_manual.csv.", "",
          "## Decisões", "",
          "- Linfangiectasia: categoria própria (o texto contém \"angiect\"; no trabalho anterior era contada "
          "como angiectasia).",
          "- Lesao_no_ileon e Congestao_da_mucosa: patologia (`outra_patologia`), como pede o guião.",
          "- Sangue conta como lesão (como no script antigo).",
          "- Estratificação por classe (guião A): 6 classes principais (angiectasia, erosão, úlcera, pólipo, "
          "eritema, xantelasma) + \"outras patologias\" (as outras 6 categorias com lesão = 1).",
          "- Não-lesão e qualidade ficam fora dos conjuntos de lesão (ALL-SEEDBOX, TEST-BOX, CAL, VAL-EX).", ""]
    Path(caminho).parent.mkdir(parents=True, exist_ok=True)
    Path(caminho).write_text("\n".join(L), encoding="utf-8")
    return True


# ------------------------------------------------------------------------------------------------
def proposta(args):
    anterior, copia = None, None
    if Path(args.revisao).exists():
        if not args.forcar:
            sys.exit(f"{args.revisao} já existe e pode ter a tua revisão: não a substituo.\n"
                     f"Para gerar a folha outra vez (guarda uma cópia e mantém o que já reviste): --forcar")
        anterior = ler_revisao(args.revisao)
        rp = Path(args.revisao)
        copia = rp.with_name(rp.stem + "_anterior" + rp.suffix)
        shutil.copy2(rp, copia)
    jan = ler_janelas(args.janelas)
    todas, n_antigas = descricoes_manifest(args.manifest)
    cont = contagens_por_descricao(jan)
    nas_janelas = {d for ds in jan["descs"] for d in ds}
    descs = sorted(set(todas) | nas_janelas)

    linhas = []
    for d in descs:
        cat, termos, motivos, antigo = classificar(d)
        c = cont.get(d, {})
        linhas.append(dict(descricao=d, categoria=cat, revisto="", notas="", categoria_auto=cat,
                           rever="sim" if motivos else "", motivo=motivos, termos=termos, regra_antiga=antigo,
                           n_janelas=c.get("n_janelas", 0), n_treino=c.get("n_treino", 0),
                           n_teste=c.get("n_teste", 0), n_com_caixa=c.get("n_com_caixa", 0),
                           n_caixa_seed=c.get("n_caixa_seed", 0), exemplo=c.get("exemplo", "")))
    herdadas = 0
    if anterior is not None and "descricao" in anterior.columns:
        ant = {r["descricao"]: r for r in anterior.to_dict("records")}
        for lin in linhas:
            a = ant.get(lin["descricao"])
            if not a:
                continue
            cat_a = cat_norm(a.get("categoria", ""))
            revisto = norm(a.get("revisto", "")) in SIM
            if revisto or cat_a != a.get("categoria_auto", "") or a.get("notas", ""):
                herdadas += 1
                if cat_a:
                    lin["categoria"] = cat_a
                lin["revisto"] = "sim" if revisto else ""
                lin["notas"] = a.get("notas", "")
                if lin["categoria"] != lin["categoria_auto"]:
                    lin["motivo"] = "; ".join(x for x in [lin["motivo"], f"ficou a tua escolha ({lin['categoria']}); "
                                                          f"a regra propõe {lin['categoria_auto'] or 'nada'}"] if x)
                    lin["rever"] = "sim"
    df = pd.DataFrame(linhas)
    df = df.assign(_o=df["categoria_auto"].map(ORDEM), _n=df["descricao"].map(norm)) \
           .sort_values(["_o", "_n"]).drop(columns=["_o", "_n"]).reset_index(drop=True)
    df.insert(0, "id", [f"D{i:04d}" for i in range(1, len(df) + 1)])
    Path(args.revisao).parent.mkdir(parents=True, exist_ok=True)
    gravar_folha(df, args.revisao)
    novo_md = criterios_md(args.criterios)

    mapa = dict(zip(df["descricao"], df["categoria"]))
    manual = ler_manual(args.manual)
    n_manual = aplicar_janelas(jan, mapa, manual)
    antigo = {d: rotulo_antigo(d) for d in descs}
    jan["lesao_antiga"] = ["1" if any(antigo[d] == "lesao" for d in ds) else ("0" if ds else "")
                           for ds in jan["descs"]]

    L = ["Mapa de rótulos — PROPOSTA AUTOMÁTICA (passo 0.3, antes da revisão à mão)"]
    sec = lambda titulo: L.append(f"\n== {titulo} " + "=" * max(0, 64 - len(titulo)))
    sec("Descrições")
    L.append(f"  descrições distintas no dataset: {len(descs)}   (diagnóstico antigo: 747)")
    L.append(f"  contadas à moda do script antigo (sem tirar pontos e espaços do fim do texto): {n_antigas}")
    L.append(f"  usadas por janelas: {len(nas_janelas)} · sem janela (imagem do relatório sem seed com frames): "
             f"{len(descs) - len(nas_janelas)}")
    L.append(f"  a rever (linhas amarelas na folha): {int((df['rever'] == 'sim').sum())} · sem regra: "
             f"{int((df['categoria_auto'] == '').sum())} · sem janela nenhuma (podes saltar): "
             f"{int((df['n_janelas'] == 0).sum())}")
    if anterior is not None:
        L.append(f"  folha anterior guardada em {copia}; linhas que já tinhas revisto ou mudado e que passaram "
                 f"para a folha nova: {herdadas}")
    mot = Counter(m.split(":")[0].split(" (")[0].split(" —")[0].split(";")[0]
                  for ms in df["motivo"] if ms for m in ms.split("; "))
    for k, v in mot.most_common():
        L.append(f"    {v:>5}  {k}")
    sec("Tabela por categoria (estado atual da folha)")
    bloco_tabela(L, jan, mapa)
    sec("Comparação com os números do diagnóstico antigo (referência, não meta)")
    bloco_conjuntos(L, jan, "lesao", "regra nova (estado atual da folha):")
    L.append("")
    bloco_conjuntos(L, jan, "lesao_antiga", "regra antiga (build_dataset.py, lesão = qualquer palavra da lista PAT_KW):")
    sec("Classes para estratificar (6 principais + outras)")
    bloco_classes(L, jan)
    sec("Colisões conhecidas (guião)")
    for d in descs:
        if re.search(r"linfang|lymphang|lesao no ileon|congestao|bulbite", norm(d)):
            r = df[df["descricao"] == d].iloc[0]
            L.append(f"  {d}: {r['categoria_auto'] or '(sem regra)'} · regra antiga: {r['regra_antiga']}"
                     f" · {r['n_janelas']} janelas")
    sec("Lesão segundo a regra nova ≠ regra antiga (por descrição)")
    dif = df[(df["categoria_auto"] != "") &
             (df["categoria_auto"].map(lambda c: LESAO.get(c, -1)) != (df["regra_antiga"] == "lesao").astype(int))]
    L.append(f"  {len(dif)} descrições")
    for r in dif.sort_values("n_janelas", ascending=False).itertuples(index=False):
        L.append(f"    {r.n_janelas:>4} jan.  {r.descricao}: nova {r.categoria_auto} (lesão {LESAO[r.categoria_auto]})"
                 f" · antiga {r.regra_antiga}")
    sec("Descrições sem regra (escolher à mão)")
    for r in df[df["categoria_auto"] == ""].sort_values("n_janelas", ascending=False).itertuples(index=False):
        L.append(f"    {r.n_janelas:>4} jan.  {r.descricao}")
    sec("Janelas com mais de uma descrição")
    bloco_multiplas(L, jan, mapa)
    if n_manual:
        L.append(f"  categorias impostas à mão ({args.manual}): {n_manual} janelas")
    sec("Ficheiros")
    L.append(f"  {args.revisao}  ← folha para rever à mão ({len(df)} linhas)")
    L.append(f"  {args.criterios}" + ("  (rascunho novo)" if novo_md else "  (já existia: não mexi)"))
    L.append("  Depois da revisão: python scripts/05_rotulos.py --aplicar")
    escrever(L, args.resumo_proposta)


def aplicar(args):
    rev = ler_revisao(args.revisao)
    erros = []
    for c in ["descricao", "categoria"]:
        if c not in rev.columns:
            erros.append(f"  falta a coluna '{c}' na folha {args.revisao}")
    if erros:
        sys.exit("\n".join(erros))
    for c in ["id", "revisto", "notas", "categoria_auto"]:
        if c not in rev.columns:
            rev[c] = ""
    rev["categoria"] = rev["categoria"].map(cat_norm)
    rev = rev[rev["descricao"] != ""].copy()

    jan = ler_janelas(args.janelas)
    todas = set(descricoes_manifest(args.manifest)[0])
    nas_janelas = {d for ds in jan["descs"] for d in ds}
    rep = rev[rev["descricao"].duplicated(keep=False)]
    for d in sorted(set(rep["descricao"])):
        erros.append(f"  descrição repetida na folha: {d}")
    invalidas = rev[~rev["categoria"].isin(CODIGOS) & (rev["categoria"] != "")]
    for r in invalidas.itertuples(index=False):
        erros.append(f"  {r.id or '?'} '{r.descricao}': categoria '{r.categoria}' não existe")
    nas_usaveis = {d for ds in jan.loc[jan["excluida"] == 0, "descs"] for d in ds}
    vazias = rev[(rev["categoria"] == "") & rev["descricao"].isin(nas_usaveis)]
    for r in vazias.itertuples(index=False):
        erros.append(f"  {r.id or '?'} '{r.descricao}': sem categoria (é usada por janelas)")
    na_folha = set(rev["descricao"])
    for d in sorted(nas_janelas - na_folha):
        erros.append(f"  descrição de janelas que não está na folha: '{d}' (linha apagada ou texto alterado?)")
    estranhas = sorted(na_folha - todas - nas_janelas)
    for d in estranhas:
        erros.append(f"  descrição na folha que não existe no dataset: '{d}' (texto da coluna B alterado?)")

    mapa = {d: (c if c in LESAO else "") for d, c in zip(rev["descricao"], rev["categoria"])}
    manual = ler_manual(args.manual)
    n_manual = aplicar_janelas(jan, mapa, manual)
    us = jan[jan["excluida"] == 0]
    por_resolver = us[us["categoria"] == ""]
    for r in por_resolver.itertuples(index=False):
        if r.categoria_origem in ("varias_patologias", "sem_descricao"):     # o resto resolve-se na folha
            erros.append(f"  janela {r.split}/{r.caso}/{r.seed} sem categoria ({r.categoria_origem}: "
                         f"{' | '.join(r.descs) or 'sem descrição'}) → decide à mão e acrescenta-a a {args.manual}")
    n_afetadas = int((por_resolver["categoria_origem"] == "descricao_sem_categoria").sum())
    if n_afetadas:
        erros.append(f"  (por causa destas linhas da folha, {n_afetadas} janelas ainda ficam sem categoria)")
    if erros:
        msg = "A folha revista tem problemas — nada foi gravado:\n" + "\n".join(erros)
        Path(args.resumo).parent.mkdir(parents=True, exist_ok=True)
        Path(args.resumo).with_name("05_rotulos_erros.txt").write_text(msg + "\n", encoding="utf-8")
        print(msg)
        print(f"\n(lista também em {Path(args.resumo).with_name('05_rotulos_erros.txt')})")
        sys.exit(1)

    # ---------- mapa final ----------
    cont = contagens_por_descricao(jan)
    rev["lesao"] = rev["categoria"].map(lambda c: str(LESAO[c]) if c else "")
    rev["revisto"] = rev["revisto"].map(lambda v: "sim" if norm(v) in SIM else "")
    rev["alterada"] = [("sim" if a and a != c else "") for a, c in zip(rev["categoria_auto"], rev["categoria"])]
    for k in ["n_janelas", "n_treino", "n_teste", "n_com_caixa", "n_caixa_seed"]:
        rev[k] = rev["descricao"].map(lambda d: cont.get(d, {}).get(k, 0))
    rev["exemplo"] = rev["descricao"].map(lambda d: cont.get(d, {}).get("exemplo", ""))
    final = rev.assign(_o=rev["categoria"].map(ORDEM), _n=rev["descricao"].map(norm)) \
               .sort_values(["_o", "_n"])[["descricao", "categoria", "lesao", "notas", "categoria_auto", "alterada",
                                           "revisto", "n_janelas", "n_treino", "n_teste", "n_com_caixa",
                                           "n_caixa_seed", "exemplo", "id"]]
    Path(args.mapa).parent.mkdir(parents=True, exist_ok=True)
    gravar_csv(final, args.mapa)

    # ---------- tabela das janelas ----------
    jw = pd.read_csv(args.janelas, dtype=str, keep_default_na=False)
    jw = jw.drop(columns=[c for c in COLS_NOVAS_JAN if c in jw.columns])
    chave = lambda df: df["split"] + "|" + df["caso"] + "|" + df["seed"]
    novo = dict(zip(chave(jan), zip(jan["categoria"], jan["lesao"], jan["categoria_origem"])))
    kj = chave(jw)
    for i, c in enumerate(COLS_NOVAS_JAN):
        jw[c] = kj.map(lambda k: novo[k][i])
    gravar_csv(jw, args.janelas)

    # ---------- tabela dos frames (por blocos, para poupar memória) ----------
    print("A acrescentar categoria e lesao à tabela dos frames (por blocos)…")
    tmp = f"{args.tabela}.tmp"
    n_linhas, frames_tb, soma_tb = 0, set(), 0
    tb = set(chave(jan[(jan["excluida"] == 0) & (jan["lesao"] == "1") & (jan["n_box"] > 0)
                       & jan["split"].str.endswith("Teste")]))
    for i, bloco in enumerate(pd.read_csv(args.tabela, dtype=str, keep_default_na=False, chunksize=200_000)):
        bloco = bloco.drop(columns=[c for c in COLS_NOVAS_FR if c in bloco.columns])
        kb = chave(bloco)
        bloco["categoria"] = kb.map(lambda k: novo.get(k, ("", "", ""))[0])
        bloco["lesao"] = kb.map(lambda k: novo.get(k, ("", "", ""))[1])
        sel = kb.isin(tb) & bloco["box_src"].isin(["medico_seed", "medico_vizinho"])
        soma_tb += int(sel.sum())
        frames_tb |= set(zip(bloco.loc[sel, "caso"], bloco.loc[sel, "frame"]))
        bloco.to_csv(tmp, index=False, mode="w" if i == 0 else "a", header=(i == 0))
        n_linhas += len(bloco)
        print(f"  {n_linhas} linhas")
    os.replace(tmp, args.tabela)

    # ---------- resumo ----------
    antigo = {d: rotulo_antigo(d) for d in mapa}
    jan["lesao_antiga"] = ["1" if any(antigo.get(d) == "lesao" for d in ds) else ("0" if ds else "")
                           for ds in jan["descs"]]
    L = ["Mapa de rótulos — FINAL (passo 0.3, depois da revisão à mão)"]
    sec = lambda titulo: L.append(f"\n== {titulo} " + "=" * max(0, 64 - len(titulo)))
    sec("Revisão")
    n_rev = int((rev["revisto"] == "sim").sum())
    L.append(f"  descrições na folha: {len(rev)}   (diagnóstico antigo: 747) · revistas (coluna revisto = sim): {n_rev}"
             + ("   ← FALTA REVER " + str(len(rev) - n_rev) if n_rev < len(rev) else "   (todas)"))
    alt = final[final["alterada"] == "sim"]
    L.append(f"  categorias mudadas à mão em relação à proposta automática: {len(alt)}"
             f" · que eram 'sem regra' e foram escolhidas: {int((final['categoria_auto'] == '').sum())}")
    sec("Tabela por categoria (final) — a 'Tabela 1' com as correções")
    bloco_tabela(L, jan, mapa)
    sec("Conjuntos (pré-visualização; os oficiais congelam-se no passo 0.6) e diagnóstico antigo")
    bloco_conjuntos(L, jan, "lesao", "regra final:")
    L.append(f"  frames com caixa no TEST-BOX de lesão: {soma_tb} (frames distintos: {len(frames_tb)})"
             f"   (diagnóstico antigo, todas as 459: 1 867)")
    L.append("")
    bloco_conjuntos(L, jan, "lesao_antiga", "regra antiga (build_dataset.py), para comparar:")
    sec("Classes para estratificar (6 principais + outras)")
    bloco_classes(L, jan)
    sec("Correções à proposta automática (declarar na tese)")
    for r in alt.sort_values("n_janelas", ascending=False).itertuples(index=False):
        L.append(f"    {r.n_janelas:>4} jan.  {r.descricao}: {r.categoria_auto} → {r.categoria}"
                 + (f"  [{r.notas}]" if r.notas else ""))
    esc = final[final["categoria_auto"] == ""]
    if len(esc):
        L.append("  sem regra, escolhidas à mão:")
        for r in esc.sort_values("n_janelas", ascending=False).itertuples(index=False):
            L.append(f"    {r.n_janelas:>4} jan.  {r.descricao}: {r.categoria}" + (f"  [{r.notas}]" if r.notas else ""))
    sec("Janelas com mais de uma descrição")
    bloco_multiplas(L, jan, mapa)
    L.append(f"  categorias impostas à mão ({args.manual}): {n_manual} janelas")
    for k, (c, mot) in manual.items():
        L.append(f"    {k[0]}/{k[1]}/{k[2]} → {c}" + (f" ({mot})" if mot else ""))
    sec("Ficheiros")
    L.append(f"  {args.mapa} ({len(final)} linhas)")
    L.append(f"  {args.janelas} (+ {', '.join(COLS_NOVAS_JAN)})")
    L.append(f"  {args.tabela} (+ {', '.join(COLS_NOVAS_FR)}; {n_linhas} linhas)")
    exemplo = jan[jan["seed_tem_box"] == 1].head(3)
    L.append("  exemplos:")
    for r in exemplo.itertuples(index=False):
        L.append(f"    {r.split}/{r.caso}/{r.seed}: {' | '.join(r.descs)} → {r.categoria} (lesão {r.lesao})")
    escrever(L, args.resumo)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--aplicar", action="store_true", help="aplica a folha revista (2.ª vez)")
    ap.add_argument("--forcar", action="store_true", help="na proposta: substitui uma folha de revisão existente")
    ap.add_argument("--manifest", default="data/derived/manifest.csv")
    ap.add_argument("--janelas", default="data/derived/janelas_v1.csv")
    ap.add_argument("--tabela", default="data/derived/annotations_v1.csv")
    ap.add_argument("--revisao", default="docs/label_map_revisao.xlsx")
    ap.add_argument("--manual", default="docs/rotulos_janelas_manual.csv")
    ap.add_argument("--criterios", default="docs/criterios_categorias.md")
    ap.add_argument("--mapa", default="data/derived/label_map.csv")
    ap.add_argument("--resumo-proposta", dest="resumo_proposta", default="outputs/05_rotulos_proposta.txt")
    ap.add_argument("--resumo", default="outputs/05_rotulos_resumo.txt")
    args = ap.parse_args()
    aplicar(args) if args.aplicar else proposta(args)


if __name__ == "__main__":
    main()
